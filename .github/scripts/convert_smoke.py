"""CI 用: [convert] を入れた環境で、サンプルの Word を harness convert する.

fieldcheck の読み込み（Markdown の表）で元の値に戻るかも確かめる.

    python .github/scripts/convert_smoke.py 作業フォルダ

依存の衝突（pip install ".[convert]" の ResolutionImpossible）や、doc2md の出力仕様の変化をこの先も検出するため。
"""

from __future__ import annotations

import shutil
import subprocess  # nosec B404 — CI で harness を動かすため
import sys
from pathlib import Path

from docx import Document

from field_validation.loaders import open_source

ROOT = Path(__file__).resolve().parents[2]
CELLS = {"項目": "パス", "値": "C:\\data\\x", "備考": "a|b <b>"}


def main(work: Path) -> int:
    if work.exists():
        shutil.rmtree(work)
    for d in ["config", "docs", "inputs", "reviews"]:
        shutil.copytree(ROOT / d, work / d)
    (work / "data").mkdir()
    doc = Document()
    doc.add_heading("サンプル API", 1)
    doc.add_paragraph("CI の確認用の架空の設計書です。")
    table = doc.add_table(rows=2, cols=len(CELLS))
    for i, (k, v) in enumerate(CELLS.items()):
        table.cell(0, i).text = k
        table.cell(1, i).text = v
    src = work / "sample.docx"
    doc.save(src)
    cmd = [sys.executable, "-m", "review_harness", "--root", str(work), "convert", str(src), "--type", "api"]
    rc = subprocess.run(cmd, check=False).returncode  # nosec B603 — 引数は固定
    if rc != 0:
        print(f"harness convert が終了コード {rc}", file=sys.stderr)
        return 1
    md = work / "inputs/api/sample.md"
    print(md.read_text(encoding="utf-8"))
    records = [r for _, r in open_source(md).records]
    if records != [CELLS]:
        print(f"読み戻した値が違います: {records} != {[CELLS]}", file=sys.stderr)
        return 1
    print("OK: harness convert と、変換結果の読み戻し")
    return 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1])))
